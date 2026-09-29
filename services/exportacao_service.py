"""Exportacao da planilha de controle (.xlsx) -- versao reduzida, so as 5
primeiras colunas do formato usado nas abas diarias da planilha de clientes
(ex: '24_09'): CLIENTE, FREQUENCIA, ERROS GERADOS, CONTROLM / MANUAL,
EXECUTADOS. Os dados vem da mesma leitura (validacoes) que alimenta a tela
de Monitoramento, entao nunca diverge do que esta sendo mostrado la."""
import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

CABECALHO = [
    "CLIENTE",
    "FREQUENCIA",
    "ERROS GERADOS",
    "CONTROLM / MANUAL",
    "EXECUTADOS",
]

_LARGURAS = [26, 12, 14, 16, 16]


def gerar_planilha_controle(validacoes, clientes):
    """validacoes: lista combinada (bad + execucao), igual a exibida no
    Monitoramento -- ja filtrada pelo chamador (selecao manual e/ou
    frequencia/horario) antes de chegar aqui. clientes: dict id->cadastro
    (ClienteModel.todos()), usado so pra pegar o CONTROLM/MANUAL cadastrado
    (nao vem em `validacoes`, que so tem o que muda por varredura)."""
    clientes_por_nome = {c["nome"]: c for c in clientes.values()}

    workbook = openpyxl.Workbook()
    planilha = workbook.active
    planilha.title = "Monitoramento"

    planilha.append(CABECALHO)

    for celula in planilha[1]:
        celula.font = Font(bold=True)
        celula.alignment = Alignment(wrap_text=True, vertical="center")

    for indice, largura in enumerate(_LARGURAS, start=1):
        planilha.column_dimensions[get_column_letter(indice)].width = largura

    for item in validacoes:
        cliente = clientes_por_nome.get(item["cliente"], {})

        erros_gerados = "BAD GERADO" if item["bad_status"] == "bad" else "-"
        executados = "OK" if item["exec_status"] == "ok" else "PENDENTE ARQUIVO"

        planilha.append([
            item["cliente"],
            item["frequencia_verificacao"],
            erros_gerados,
            cliente.get("controlm_manual", "-"),
            executados,
        ])

    for linha_celulas in planilha.iter_rows(min_row=2):
        for celula in linha_celulas:
            celula.alignment = Alignment(vertical="top")

    planilha.freeze_panes = "A2"
    planilha.auto_filter.ref = f"A1:{get_column_letter(len(CABECALHO))}{planilha.max_row}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return buffer


def nome_arquivo_exportacao():
    return f"monitoramento_{datetime.now().strftime('%d-%m-%Y_%H%M')}.xlsx"
