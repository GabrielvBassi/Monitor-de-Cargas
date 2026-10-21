"""Preenchimento incremental da aba diaria DENTRO da propria planilha de
clientes (ex: '24_09', '25_09'...) -- mesmo formato que a area ja mantem na
mao todo dia (12 colunas: CLIENTE, FREQUENCIA, ERROS GERADOS, CONTROLM /
MANUAL, EXECUTADOS, DIRETORIOS, DIRETORIOS BADS, ULTIMO ARQUIVO, HORARIO
EXECUCAO, ACOMPANHAMENTO DIARIO/SEMANAL/MENSAL).

Pensado pra ir preenchendo aos poucos ao longo do dia, por leva de horario
de processamento: cada chamada so grava os clientes passados em
`validacoes` (o chamador filtra antes, ex: so quem processa as 9h/10h/11h) e
NAO mexe nas linhas ja gravadas de levas anteriores -- clientes repetidos
(mesmo nome ja gravado antes nessa aba) tem a linha ATUALIZADA no lugar,
nunca duplicada.

Ao contrario de exportacao_service.py (que gera um .xlsx separado pra
baixar, reduzido a 5 colunas), aqui escreve direto no arquivo configurado em
config.CLIENTES_XLSX_PATH -- o mesmo arquivo que alimenta o cadastro de
clientes do app."""
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

import config
from models.cliente_model import ClienteModel, _normalizar

CABECALHO = [
    "CLIENTE",
    "FREQUENCIA",
    "ERROS GERADOS",
    "CONTROLM / MANUAL",
    "EXECUTADOS",
    "DIRETÓRIOS",
    "DIRETÓRIOS BADS",
    "ÚLTIMO ARQUIVO",
    "HORÁRIO EXECUÇÃO",
    "ACOMPANHAMENTO DIÁRIO",
    "ACOMPANHAMENTO SEMANAL",
    "ACOMPANHAMENTO MENSAL",
]

_LARGURAS = [26, 12, 14, 16, 16, 45, 45, 14, 14, 18, 18, 18]


def nome_aba_do_dia(data=None):
    """Mesmo padrao das abas que ja existem na planilha (ex: '24_09')."""
    return (data or date.today()).strftime("%d_%m")


def _formatar_pastas(pastas):
    if not pastas:
        return "-"

    return "\n".join(pastas)


def _linha_valores(item, clientes_por_nome, acompanhamento_por_nome):
    cliente = clientes_por_nome.get(item["cliente"], {})
    acompanhamento = acompanhamento_por_nome.get(_normalizar(item["cliente"]), {})

    erros_gerados = "BAD GERADO" if item["bad_status"] == "bad" else "-"
    executados = "OK" if item["exec_status"] == "ok" else "PENDENTE ARQUIVO"
    ultimo_arquivo = (
        item["exec_data_arquivo"].strftime("%d/%m/%Y")
        if item.get("exec_data_arquivo")
        else "-"
    )

    return [
        item["cliente"],
        item["frequencia_verificacao"],
        erros_gerados,
        cliente.get("controlm_manual", "-"),
        executados,
        _formatar_pastas(item.get("exec_pastas")),
        _formatar_pastas(item.get("bad_pastas")),
        ultimo_arquivo,
        item.get("horario_execucao", "-"),
        acompanhamento.get("diario", "-") or "-",
        acompanhamento.get("semanal", "-") or "-",
        acompanhamento.get("mensal", "-") or "-",
    ]


def atualizar_aba_diaria(validacoes, clientes, data=None):
    """Grava, na aba do dia (criando se ainda nao existir), so as linhas dos
    clientes presentes em `validacoes` -- o chamador filtra essa lista antes
    (ex: so os clientes do horario selecionado na tela).

    Clientes que ja tinham linha na aba (de uma leva anterior no mesmo dia)
    e NAO estao em `validacoes` desta chamada ficam como estavam, intactos.
    Clientes que aparecem de novo (mesmo nome ja gravado antes) tem a linha
    deles SOBRESCRITA com os dados de agora, nunca duplicada.

    Levanta PermissionError se o arquivo estiver aberto/bloqueado por outro
    programa (ex: alguem com o Excel aberto nele) -- nesse caso nada foi
    alterado no arquivo original.

    Retorna (nome_aba, quantidade_criada, quantidade_atualizada)."""
    nome_aba = nome_aba_do_dia(data)

    # data_only=False: preserva as formulas das outras abas. Se abrisse com
    # data_only=True, ao salvar de volta as formulas virariam o ultimo
    # valor calculado (perda permanente).
    workbook = openpyxl.load_workbook(config.CLIENTES_XLSX_PATH, data_only=False)

    if nome_aba in workbook.sheetnames:
        planilha = workbook[nome_aba]
    else:
        planilha = workbook.create_sheet(nome_aba)
        planilha.append(CABECALHO)

        for celula in planilha[1]:
            celula.font = Font(bold=True)
            celula.alignment = Alignment(wrap_text=True, vertical="center")

        for indice, largura in enumerate(_LARGURAS, start=1):
            planilha.column_dimensions[get_column_letter(indice)].width = largura

    # Mapeia nome do cliente -> numero da linha ja existente na aba, pra
    # atualizar em vez de duplicar quando o mesmo cliente aparece de novo
    # numa leva posterior.
    linha_por_cliente = {}

    for indice_linha in range(2, planilha.max_row + 1):
        nome_celula = planilha.cell(row=indice_linha, column=1).value

        if nome_celula:
            linha_por_cliente[str(nome_celula).strip()] = indice_linha

    clientes_por_nome = {c["nome"]: c for c in clientes.values()}
    acompanhamento_por_nome = ClienteModel.acompanhamento()

    criados = 0
    atualizados = 0

    for item in validacoes:
        valores = _linha_valores(item, clientes_por_nome, acompanhamento_por_nome)
        indice_linha = linha_por_cliente.get(item["cliente"])

        if indice_linha:
            for indice_coluna, valor in enumerate(valores, start=1):
                planilha.cell(row=indice_linha, column=indice_coluna, value=valor)

            atualizados += 1
        else:
            planilha.append(valores)
            linha_por_cliente[item["cliente"]] = planilha.max_row
            criados += 1

    for linha_celulas in planilha.iter_rows(min_row=2):
        for celula in linha_celulas:
            celula.alignment = Alignment(wrap_text=True, vertical="top")

    planilha.freeze_panes = "A2"
    planilha.auto_filter.ref = f"A1:{get_column_letter(len(CABECALHO))}{planilha.max_row}"

    try:
        workbook.save(config.CLIENTES_XLSX_PATH)
    except PermissionError as exc:
        raise PermissionError(
            "Nao foi possivel salvar: a planilha parece estar aberta "
            "(por voce ou por outra pessoa, no Excel). Feche o arquivo e "
            "tente de novo -- nada foi alterado."
        ) from exc

    return nome_aba, criados, atualizados
